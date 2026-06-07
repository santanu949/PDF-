import requests
import time
import random
import os

time.sleep(5)

def create_dummy_jpg():
    try:
        from PIL import Image
        Image.new('RGB', (100, 100), color='red').save('test.jpg')
    except ImportError:
        with open('test.jpg', 'wb') as f:
            f.write(b'\xff\xd8\xff\xe0\x00\x10JFIF\x00\x01\x01\x01\x00H\x00H\x00\x00\xff\xdb\x00C\x00\x05\x03')

def create_dummy_docx():
    try:
        import docx
        docx.Document().save('test.docx')
    except ImportError:
        with open('test.docx', 'wb') as f:
            f.write(b'PK\x03\x04')

def create_dummy_xlsx():
    try:
        import openpyxl
        openpyxl.Workbook().save('test.xlsx')
    except ImportError:
        with open('test.xlsx', 'wb') as f:
            f.write(b'PK\x03\x04')

def create_dummy_pptx():
    try:
        from pptx import Presentation
        Presentation().save('test.pptx')
    except ImportError:
        with open('test.pptx', 'wb') as f:
            f.write(b'PK\x03\x04')

create_dummy_jpg()
create_dummy_docx()
create_dummy_xlsx()
create_dummy_pptx()

BASE_URL = "http://localhost:8000/api/v1"
email = f"test_full_{random.randint(1000, 99999)}@example.com"
password = "password123"
headers = {}

results = []

def record(test_name, status, reason="", extra=""):
    print(f"{test_name}: {status} {reason} {extra}")
    results.append({"name": test_name, "status": status, "reason": reason, "extra": extra})

def poll_job(job_id):
    while True:
        r = requests.get(f"{BASE_URL}/jobs/{job_id}/status", headers=headers)
        if r.status_code != 200:
            return "failed", r.text
        st = r.json().get("status")
        if st in ["completed", "failed"]:
            return st, r.json().get("error_message", "")
        time.sleep(2)

def run_test(name, endpoint, payload=None, files=None):
    try:
        if files:
            r = requests.post(f"{BASE_URL}{endpoint}", headers=headers, files=files)
        else:
            r = requests.post(f"{BASE_URL}{endpoint}", headers=headers, json=payload)
        
        if r.status_code != 200:
            record(name, "FAILED", f"HTTP {r.status_code}", r.text)
            return None
        
        data = r.json()
        if "job_id" in data:
            st, err = poll_job(data["job_id"])
            if st == "completed":
                rd = requests.get(f"{BASE_URL}/jobs/{data['job_id']}/download", headers=headers)
                if rd.status_code == 200:
                    record(name, "PASSED")
                    return data
                else:
                    record(name, "FAILED", "Download failed", f"HTTP {rd.status_code}: {rd.text}")
                    return data
            else:
                record(name, "FAILED", "Job failed", err)
                return data
        else:
            record(name, "PASSED")
            return data
    except Exception as e:
        record(name, "FAILED", "Exception", str(e))
        return None

# Auth
r1 = requests.post(f"{BASE_URL}/auth/register", json={"email": email, "password": password})
if r1.status_code == 200: record("Register", "PASSED")
else: record("Register", "FAILED", f"HTTP {r1.status_code}", r1.text)

r2 = requests.post(f"{BASE_URL}/auth/login", json={"email": email, "password": password})
if r2.status_code == 200:
    record("Login", "PASSED")
    token = r2.json().get("access_token")
    headers["Authorization"] = f"Bearer {token}"
else:
    record("Login", "FAILED", f"HTTP {r2.status_code}", r2.text)

r3 = requests.get(f"{BASE_URL}/auth/me", headers=headers)
if r3.status_code == 200: record("JWT Protected Route", "PASSED")
else: record("JWT Protected Route", "FAILED", f"HTTP {r3.status_code}", r3.text)

# Uploads
pdf_file_id = None
pdf_file_id2 = None
if headers:
    r_pdf = requests.post(f"{BASE_URL}/files/upload", headers=headers, files={'file': open('test1.pdf', 'rb')})
    if r_pdf.status_code == 200:
        record("Upload PDF 1", "PASSED")
        pdf_file_id = r_pdf.json()["file_id"]
    else:
        record("Upload PDF 1", "FAILED", f"HTTP {r_pdf.status_code}", r_pdf.text)
        
    r_pdf2 = requests.post(f"{BASE_URL}/files/upload", headers=headers, files={'file': open('test2.pdf', 'rb')})
    if r_pdf2.status_code == 200:
        record("Upload PDF 2", "PASSED")
        pdf_file_id2 = r_pdf2.json()["file_id"]
    else:
        record("Upload PDF 2", "FAILED", f"HTTP {r_pdf2.status_code}", r_pdf2.text)

jpg_id, docx_id, xlsx_id, pptx_id = None, None, None, None
if headers:
    r_jpg = requests.post(f"{BASE_URL}/files/upload", headers=headers, files={'file': open('test.jpg', 'rb')})
    if r_jpg.status_code == 200: jpg_id = r_jpg.json()["file_id"]

    r_docx = requests.post(f"{BASE_URL}/files/upload", headers=headers, files={'file': open('test.docx', 'rb')})
    if r_docx.status_code == 200: docx_id = r_docx.json()["file_id"]

    r_xlsx = requests.post(f"{BASE_URL}/files/upload", headers=headers, files={'file': open('test.xlsx', 'rb')})
    if r_xlsx.status_code == 200: xlsx_id = r_xlsx.json()["file_id"]

    r_pptx = requests.post(f"{BASE_URL}/files/upload", headers=headers, files={'file': open('test.pptx', 'rb')})
    if r_pptx.status_code == 200: pptx_id = r_pptx.json()["file_id"]

if pdf_file_id:
    # PDF Tools
    if pdf_file_id2:
        run_test("Merge PDF", "/merge/", payload={"file_ids": [pdf_file_id, pdf_file_id2]})
    run_test("Split PDF", "/split/", payload={"file_id": pdf_file_id, "pages": "1"})
    run_test("Compress PDF", "/tools/compress-pdf", payload={"file_id": pdf_file_id, "compression_percent": 50})
    run_test("Rotate PDF", "/rotate/", payload={"file_id": pdf_file_id, "rotation": 90})
    run_test("Watermark PDF", "/watermark/", payload={"file_id": pdf_file_id, "text": "Draft"})
    run_test("Add Page Numbers", "/page-numbers/", payload={"file_id": pdf_file_id, "position": "bottom-right"})
    run_test("Organize PDF", "/tools/organize-pdf", payload={"file_id": pdf_file_id, "page_order": [1]})
    run_test("Protect PDF", "/protect/", payload={"file_id": pdf_file_id, "password": "pass"})
    run_test("Unlock PDF", "/unlock/", payload={"file_id": pdf_file_id, "password": "pass"})
    run_test("Sign PDF", "/tools/sign-pdf", payload={"pdf_file_id": pdf_file_id, "mode": "typed", "signature_text": "Approved"})

    # Conversion PDF ->
    run_test("PDF to JPG", "/pdf-to-jpg/", payload={"file_id": pdf_file_id})
    run_test("PDF to Word", "/tools/pdf-to-word", payload={"file_id": pdf_file_id})
    run_test("PDF to Excel", "/tools/pdf-to-excel", payload={"file_id": pdf_file_id})
    
    # OCR & AI
    run_test("OCR PDF", "/tools/ocr", payload={"file_id": pdf_file_id, "language": "eng"})
    run_test("AI Summarize", "/ai-summarise/", payload={"file_id": pdf_file_id})
    run_test("AI Translate", "/ai-translate/", payload={"file_id": pdf_file_id, "target_language": "Spanish"})
    run_test("AI Rewrite", "/ai-rewrite/", payload={"file_id": pdf_file_id, "tone": "professional"})

if jpg_id: run_test("JPG to PDF", "/jpg-to-pdf/", payload={"file_ids": [jpg_id]})
if docx_id: run_test("Word to PDF", "/tools/word-to-pdf", payload={"file_id": docx_id})
if xlsx_id: run_test("Excel to PDF", "/tools/excel-to-pdf", payload={"file_id": xlsx_id})
if pptx_id: run_test("PPT to PDF", "/tools/ppt-to-pdf", payload={"file_id": pptx_id})

# QR Tools
run_test("QR Generator", "/tools/generate-qr", payload={"url": "https://example.com"})
if pdf_file_id:
    run_test("QR PDF", "/qr-pdf/", payload={"file_id": pdf_file_id, "url": "https://example.com"})

print("\n--- FINAL REPORT ---")
for r in results:
    print(f"{r['name']}: {r['status']} | {r['reason']} | {r['extra']}")
